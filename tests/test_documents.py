import io
import zipfile
import unittest
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from backend import documents


class DocumentValidationTests(unittest.TestCase):
    def test_upload_accepts_pdf_signature_and_rejects_unapproved_types(self):
        name, extension = documents._validate_payload(
            "policy.pdf", b"%PDF-1.7\ncontent"
        )
        self.assertEqual((name, extension), ("policy.pdf", ".pdf"))

        with self.assertRaisesRegex(documents.DocumentError, "Only PDF"):
            documents._validate_payload("data.xlsx", b"PK\x03\x04")

        with self.assertRaisesRegex(documents.DocumentError, "valid PDF"):
            documents._validate_payload("broken.pdf", b"not a pdf")

    def test_word_upload_requires_docx_package_parts(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("word/document.xml", "<document/>")

        self.assertEqual(
            documents._validate_payload("policy.docx", stream.getvalue()),
            ("policy.docx", ".docx"),
        )

    def test_chunks_preserve_page_and_overlap(self):
        source = " ".join(f"word{index}" for index in range(500))
        chunks = documents._split_text(source, page=7, section="Risk")

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk["page"] == 7 for chunk in chunks))
        self.assertTrue(all(chunk["section"] == "Risk" for chunk in chunks))
        self.assertLessEqual(
            max(len(chunk["text"]) for chunk in chunks),
            documents.CHUNK_CHARACTERS,
        )
        self.assertTrue(
            set(chunks[0]["text"].split()) & set(chunks[1]["text"].split())
        )

    def test_docx_extraction_preserves_section_heading(self):
        from docx import Document

        stream = io.BytesIO()
        document = Document()
        document.add_heading("Travel policy", level=1)
        document.add_paragraph("Manager approval is required.")
        document.save(stream)

        chunks = documents.extract_document(".docx", stream.getvalue())
        self.assertEqual(chunks[0]["section"], "Travel policy")
        self.assertIn(
            "Manager approval",
            " ".join(chunk["text"] for chunk in chunks),
        )
        self.assertTrue(all(chunk["section"] == "Travel policy" for chunk in chunks))

    def test_scanned_pdf_without_extractable_text_is_rejected_for_ocr(self):
        from pypdf import PdfWriter

        stream = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(stream)

        with self.assertRaisesRegex(documents.DocumentError, "OCR"):
            documents.extract_document(".pdf", stream.getvalue())


class DocumentPermissionTests(unittest.TestCase):
    def test_access_filter_scopes_department_user_to_owner_or_department(self):
        clause, parameters = documents._access_filter(
            {
                "username": "it_head",
                "role": "Department Head",
                "departments": ["IT"],
            }
        )

        self.assertIn("d.owner = ?", clause)
        self.assertIn("d.department IN (?)", clause)
        self.assertEqual(parameters, ["it_head", "IT"])

    def test_unrestricted_user_sees_only_owned_or_shared_documents(self):
        clause, parameters = documents._access_filter(
            {"username": "finance", "role": "Finance Manager", "departments": None}
        )

        self.assertIn("d.department IS NOT NULL", clause)
        self.assertEqual(parameters, ["finance"])

    def test_empty_department_scope_cannot_search_shared_documents(self):
        clause, parameters = documents._access_filter(
            {"username": "restricted", "role": "Department Head", "departments": []}
        )

        self.assertEqual(clause, "d.owner = ?")
        self.assertEqual(parameters, ["restricted"])

    def test_admin_access_and_department_sharing(self):
        self.assertEqual(
            documents._access_filter(
                {"username": "admin", "role": "Admin"}
            ),
            ("TRUE", []),
        )
        self.assertEqual(
            documents._shared_department(
                {"role": "Admin"}, "Finance"
            ),
            "Finance",
        )
        with self.assertRaisesRegex(documents.DocumentError, "cannot share"):
            documents._shared_department(
                {"role": "Analyst", "departments": None}, "Finance"
            )

    def test_search_sql_filters_accessible_documents_before_retrieval(self):
        cursor = MagicMock()
        cursor.fetchall.side_effect = [
            [(1, "doc-id", "policy.pdf", 3, None, "Approved policy text", 0.8)],
            [],
        ]
        connection = MagicMock()
        connection.__enter__.return_value = connection
        connection.cursor.return_value.__enter__.return_value = cursor

        with patch.object(
            documents.database, "connect", return_value=connection
        ):
            results = documents._retrieval_context(
                {
                    "username": "it_head",
                    "role": "Department Head",
                    "departments": ["IT"],
                },
                "travel policy",
                [0.0] * documents.EMBEDDING_DIMENSIONS,
            )

        sql, parameters = cursor.execute.call_args_list[0].args
        self.assertIn("cosine_similarity", sql)
        self.assertIn("d.department IN (?)", sql)
        self.assertEqual(parameters[1:3], ["it_head", "IT"])
        self.assertEqual(results[0]["page"], 3)


class DocumentAnswerTests(unittest.TestCase):
    @staticmethod
    def _completion(content):
        return SimpleNamespace(
            choices=[
                SimpleNamespace(
                    message=SimpleNamespace(content=content)
                )
            ]
        )

    def test_answer_requires_a_source_citation_and_passes_evidence_check(self):
        client = MagicMock()
        client.chat.completions.create.side_effect = [
            self._completion("The policy requires approval [S1]."),
            self._completion('{"supported": true}'),
        ]
        passages = [
            {
                "name": "travel.pdf",
                "page": 4,
                "section": None,
                "text": "Expenses require manager approval.",
            }
        ]

        result = documents._answer_from_passages(
            "What approval is required?", passages, client, "test-model"
        )

        self.assertEqual(result["reliability_status"], "Partially verified")
        self.assertIn("[S1]", result["answer"])
        self.assertEqual(client.chat.completions.create.call_count, 2)

    def test_answer_is_rejected_when_evidence_verifier_disagrees(self):
        client = MagicMock()
        client.chat.completions.create.side_effect = [
            self._completion("All expenses are reimbursed [S1]."),
            self._completion('{"supported": false}'),
        ]
        passages = [
            {
                "name": "travel.pdf",
                "page": 4,
                "section": None,
                "text": "Expenses require manager approval.",
            }
        ]

        result = documents._answer_from_passages(
            "Are all expenses reimbursed?", passages, client, "test-model"
        )

        self.assertEqual(result["reliability_status"], "No answer")
        self.assertNotIn("[S1]", result["answer"])

    def test_answer_rejects_uncited_or_unknown_citations(self):
        passages = [
            {
                "name": "policy.pdf",
                "page": 1,
                "section": None,
                "text": "Supported source text.",
            }
        ]
        for answer in ("Unsupported answer.", "Unsupported [S9]."):
            client = MagicMock()
            client.chat.completions.create.return_value = self._completion(answer)
            result = documents._answer_from_passages(
                "Question?", passages, client, "test-model"
            )
            self.assertEqual(result["reliability_status"], "No answer")
            self.assertEqual(client.chat.completions.create.call_count, 1)

    def test_answer_supports_orchestrator_document_citation_prefix(self):
        client = MagicMock()
        client.chat.completions.create.side_effect = [
            self._completion("Approval is required [D1]."),
            self._completion('{"supported": true}'),
        ]
        passages = [
            {
                "name": "travel.pdf",
                "document_id": "doc-1",
                "page": 4,
                "section": None,
                "text": "Expenses require manager approval.",
            }
        ]

        result = documents.answer_from_passages(
            "What approval is required?",
            passages,
            client=client,
            model="test-model",
            source_prefix="D",
        )

        self.assertEqual(result["reliability_status"], "Partially verified")
        self.assertEqual(result["sources"][0]["source"], "D1")


if __name__ == "__main__":
    unittest.main()
