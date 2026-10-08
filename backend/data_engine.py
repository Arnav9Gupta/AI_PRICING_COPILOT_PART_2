"""
data_engine.py
==============
All data logic for AI Pricing Copilot. No LLM calls live here: every number
the chatbot shows is computed in this file (this is what prevents hallucination).

The engine now serves FOUR datasets from one place:

    transactions   general finance transactions (income and expense)
    procurement    purchase orders
    payroll        HR payroll records
    sales          sales orders

Each dataset is described once in DATASET_CONFIGS (columns, which column is the
"actual" amount, which is the budget/target, which columns can be filtered or
grouped, and so on). All the functions below read that configuration, so adding
a fifth dataset later only needs a new entry in DATASET_CONFIGS.

Covers: record search, category/department/date filters, sorting, top/bottom
records, income-expense summary, budget vs actual, variance detection,
comparison, grouping, input validation and access control.
"""

import os
import hmac
import hashlib
import difflib
from collections import defaultdict
from datetime import datetime, date

from openpyxl import load_workbook


# ============================================================
# DATASET LOCATION
# ============================================================

BASE_DIR = os.path.dirname(
    os.path.dirname(os.path.abspath(__file__))
)

DATA_DIR = os.path.join(BASE_DIR, "Data")

DATE_FORMATS = [
    "%Y-%m-%d",
    "%d-%m-%Y",
    "%m-%d-%Y",
    "%d/%m/%Y",
    "%m/%d/%Y",
    "%Y/%m/%d",
    "%d-%b-%Y"
]


# ============================================================
# DATASET CONFIGURATION
# ============================================================
# Keys used below
#   id / date / amount / budget   core columns ("amount" = the ACTUAL value)
#   type_column      column that says Income or Expense per row (or None)
#   nature           "expense" or "income" for datasets without a type column
#   primary_kind     the kind of row that budget-vs-actual is measured on
#   department       column used for department-level access control (or None)
#   status_column    column holding the record status
#   filter_columns   categorical columns (exact / fuzzy match on real values)
#   text_columns     free-text columns (partial match)
#   exact_columns    columns that must NEVER be fuzzy matched (IDs)
#   group_columns    columns usable in "by ..." breakdowns ("Month" is always allowed)
#   measures         numeric columns that can be totalled / averaged / ranked
#   money_columns    numeric columns that are money (the rest are plain numbers)
#   warn_values      statuses that are included in totals but worth flagging
#   fallbacks        formulas used only if an Excel cell has no saved value

DATASET_CONFIGS = {

    "transactions": {
        "label": "Transactions",
        "description": "General finance transactions (income and expenses) "
                       "by department, category, vendor/customer and "
                       "payment status.",
        "row_meaning": "one financial transaction",
        "file": "synthetic_transactions_1000.xlsx",
        "sheet": None,
        "id": "Transaction ID",
        "date": "Transaction Date",
        "amount": "Amount",
        "budget": "Budget",
        "amount_label": "Amount",
        "budget_label": "Budget",
        "type_column": "Transaction Type",
        "nature": None,
        "primary_kind": "expense",
        "department": "Department",
        "status_column": "Payment Status",
        "filter_columns": ["Department", "Category", "Transaction Type",
                           "Payment Status", "Currency"],
        "text_columns": ["Vendor / Customer"],
        "exact_columns": [],
        "group_columns": ["Department", "Category", "Vendor / Customer",
                          "Payment Status", "Transaction Type", "Currency"],
        "default_group": "Category",
        "dashboard_groups": ["Department", "Category", "Month"],
        "measures": ["Amount", "Budget"],
        "money_columns": ["Amount", "Budget"],
        "numeric_columns": ["Amount", "Budget"],
        "date_columns": ["Transaction Date"],
        "warn_values": {},
        "fallbacks": {},
        "noun": "transaction",
        "noun_plural": "transactions",
        "top_label": "expenses",
        "budget_word": "budget",
        "actual_word": "spending",
        "examples": ["Show total expenses for IT in May 2026",
                     "Compare IT and HR spending"],
    },

    "procurement": {
        "label": "Procurement Orders",
        "description": "Purchase orders placed with suppliers: items, "
                       "warehouses, costs, delivery status and delays.",
        "row_meaning": "one purchase order",
        "file": "procurement_orders_700.xlsx",
        "sheet": "Procurement",
        "id": "PO Number",
        "date": "PO Date",
        "amount": "Total Cost",
        "budget": "Budgeted Cost",
        "amount_label": "Total Cost",
        "budget_label": "Budgeted Cost",
        "type_column": None,
        "nature": "expense",
        "primary_kind": "expense",
        "department": None,
        "status_column": "Delivery Status",
        "filter_columns": ["Supplier", "Item Category", "Item Name",
                           "Warehouse", "Delivery Status", "Currency"],
        "text_columns": [],
        "exact_columns": [],
        "group_columns": ["Supplier", "Item Category", "Item Name",
                          "Warehouse", "Delivery Status", "Currency"],
        "default_group": "Item Category",
        "dashboard_groups": ["Item Category", "Supplier", "Warehouse", "Month"],
        "measures": ["Total Cost", "Budgeted Cost", "Qty Ordered",
                     "Qty Received", "Unit Cost", "Delay Days"],
        "money_columns": ["Total Cost", "Budgeted Cost", "Unit Cost"],
        "numeric_columns": ["Qty Ordered", "Qty Received", "Unit Cost",
                            "Total Cost", "Budgeted Cost", "Delay Days"],
        "date_columns": ["PO Date", "Expected Delivery"],
        "warn_values": {"Delivery Status": ["Cancelled"]},
        "fallbacks": {
            "Total Cost": lambda r: (r["Qty Ordered"] or 0) * (r["Unit Cost"] or 0),
        },
        "noun": "purchase order",
        "noun_plural": "purchase orders",
        "top_label": "purchase orders",
        "budget_word": "budget",
        "actual_word": "spend",
        "examples": ["Which supplier has the highest total cost?",
                     "Show delayed orders for the Mumbai warehouse"],
    },

    "payroll": {
        "label": "HR Payroll",
        "description": "Monthly payroll per employee: basic salary, "
                       "allowances, bonus, deductions, gross and net pay.",
        "row_meaning": "one payroll record (one employee for one month)",
        "file": "hr_payroll_540.xlsx",
        "sheet": "Payroll",
        "id": "Payroll ID",
        "date": "Pay Date",
        "amount": "Gross Pay",
        "budget": "Budgeted Cost",
        "amount_label": "Gross Pay",
        "budget_label": "Budgeted Cost",
        "type_column": None,
        "nature": "expense",
        "primary_kind": "expense",
        "department": "Department",
        "status_column": "Payment Status",
        "filter_columns": ["Department", "Location", "Employment Type",
                           "Payment Status", "Employee ID", "Currency"],
        "text_columns": ["Employee Name", "Designation"],
        "exact_columns": ["Employee ID"],
        "group_columns": ["Department", "Location", "Employment Type",
                          "Designation", "Payment Status", "Employee Name"],
        "default_group": "Department",
        "dashboard_groups": ["Department", "Location", "Employment Type", "Month"],
        "measures": ["Gross Pay", "Net Pay", "Basic Salary", "Allowances",
                     "Bonus", "Deductions", "Budgeted Cost"],
        "money_columns": ["Gross Pay", "Net Pay", "Basic Salary", "Allowances",
                          "Bonus", "Deductions", "Budgeted Cost"],
        "numeric_columns": ["Basic Salary", "Allowances", "Bonus",
                            "Gross Pay", "Deductions", "Net Pay",
                            "Budgeted Cost"],
        "date_columns": ["Pay Date"],
        "warn_values": {},
        "fallbacks": {
            "Gross Pay": lambda r: (r["Basic Salary"] or 0)
                                   + (r["Allowances"] or 0) + (r["Bonus"] or 0),
            "Net Pay": lambda r: (r["Gross Pay"] or 0) - (r["Deductions"] or 0),
        },
        "noun": "payroll record",
        "noun_plural": "payroll records",
        "top_label": "payroll records",
        "budget_word": "budget",
        "actual_word": "payroll cost (gross pay)",
        "examples": ["What is the total net pay for Engineering?",
                     "Show bonus paid in June 2026"],
    },

    "sales": {
        "label": "Sales Orders",
        "description": "Customer sales orders by region, sales rep, product, "
                       "discount, order status and payment method.",
        "row_meaning": "one sales order",
        "file": "sales_orders_1000.xlsx",
        "sheet": "Sales Orders",
        "id": "Order ID",
        "date": "Order Date",
        "amount": "Order Value",
        "budget": "Target Value",
        "amount_label": "Order Value",
        "budget_label": "Target Value",
        "type_column": None,
        "nature": "income",
        "primary_kind": "income",
        "department": None,
        "status_column": "Order Status",
        "filter_columns": ["Region", "Sales Rep", "Customer",
                           "Product Category", "Product", "Order Status",
                           "Payment Method", "Currency"],
        "text_columns": [],
        "exact_columns": [],
        "group_columns": ["Region", "Sales Rep", "Customer",
                          "Product Category", "Product", "Order Status",
                          "Payment Method", "Currency"],
        "default_group": "Region",
        "dashboard_groups": ["Region", "Product Category", "Sales Rep", "Month"],
        "measures": ["Order Value", "Target Value", "Quantity", "Unit Price"],
        "money_columns": ["Order Value", "Target Value", "Unit Price"],
        "numeric_columns": ["Quantity", "Unit Price", "Discount %",
                            "Order Value", "Target Value"],
        "date_columns": ["Order Date"],
        "warn_values": {"Order Status": ["Cancelled", "Returned"]},
        "fallbacks": {
            "Order Value": lambda r: round(
                (r["Quantity"] or 0) * (r["Unit Price"] or 0)
                * (1 - (r["Discount %"] or 0))
            ),
        },
        "noun": "sales order",
        "noun_plural": "sales orders",
        "top_label": "sales orders",
        "budget_word": "target",
        "actual_word": "sales",
        "examples": ["Total sales by region",
                     "Which sales rep missed the target the most?"],
    },
}


# ============================================================
# SMALL HELPERS
# ============================================================

def _to_float(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _norm(text):
    """Lower-case, trimmed, underscores treated as spaces."""
    return str(text).strip().lower().replace("_", " ")


def _parse_cell_date(value):
    """Converts an Excel cell value into a datetime (or None)."""

    if isinstance(value, datetime):
        return value

    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())

    if value:
        text = str(value).strip()

        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                continue

    return None


def parse_date_input(value, end_of_day=False):
    """
    Converts a YYYY-MM-DD string (or datetime) into a datetime.
    Raises ValueError with a friendly message when invalid.
    """

    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        parsed = value
    else:
        try:
            parsed = datetime.strptime(
                str(value).strip()[:10],
                "%Y-%m-%d"
            )
        except ValueError:
            raise ValueError(
                f"'{value}' is not a valid date. Please use YYYY-MM-DD."
            )

    if end_of_day:
        parsed = parsed.replace(hour=23, minute=59, second=59)

    return parsed


def resolve_column(columns, name):
    """Case-insensitive column lookup. Returns the real column name or None."""

    if name is None:
        return None

    target = _norm(name)

    for column in columns:
        if _norm(column) == target:
            return column

    return None


# ============================================================
# LOAD DATASETS
# ============================================================

def _load_one(key, cfg):

    path = os.path.join(DATA_DIR, cfg["file"])

    if not os.path.exists(path):
        raise FileNotFoundError(f"File not found: {path}")

    workbook = load_workbook(path, data_only=True)

    sheet = (
        workbook[cfg["sheet"]]
        if cfg["sheet"] and cfg["sheet"] in workbook.sheetnames
        else workbook.active
    )

    headers = [
        str(cell.value).strip() if cell.value is not None else ""
        for cell in sheet[1]
    ]

    needed = [cfg["id"], cfg["date"], cfg["amount"], cfg["budget"]]
    missing = [c for c in needed if c not in headers]

    if missing:
        raise ValueError(
            f"Sheet '{sheet.title}' is missing column(s): {', '.join(missing)}"
        )

    rows = []

    for raw in sheet.iter_rows(min_row=2, values_only=True):

        # skip completely empty rows
        if not any(v is not None and str(v).strip() != "" for v in raw):
            continue

        record = dict(zip(headers, raw))
        record.pop("", None)

        for column in cfg["date_columns"]:
            if column in record:
                record[column] = _parse_cell_date(record[column])

        for key_, value in record.items():
            if isinstance(value, str):
                record[key_] = value.strip()

        # numbers: keep missing as None first, so formulas can fill them
        for column in cfg["numeric_columns"]:
            if column in record:
                value = record[column]
                record[column] = (
                    None if value in (None, "") else _to_float(value)
                )

        for column, formula in cfg["fallbacks"].items():
            if column in record and record[column] is None:
                try:
                    record[column] = float(formula(record))
                except (KeyError, TypeError, ValueError):
                    pass

        for column in cfg["numeric_columns"]:
            if column in record and record[column] is None:
                record[column] = 0.0

        rows.append(record)

    return rows, [h for h in headers if h]


DATASETS = {}       # key -> config (+ "columns"), only datasets that loaded
_ROWS = {}          # key -> list of row dicts
LOAD_ERRORS = {}    # key -> reason a dataset could not be loaded

for _key, _cfg in DATASET_CONFIGS.items():
    try:
        _rows, _columns = _load_one(_key, _cfg)
        _ROWS[_key] = _rows
        DATASETS[_key] = {**_cfg, "columns": _columns}
    except Exception as exc:                       # keep the app alive
        LOAD_ERRORS[_key] = f"{type(exc).__name__}: {exc}"

if not DATASETS:
    raise FileNotFoundError(
        "No dataset could be loaded from "
        f"{DATA_DIR}. Details: {LOAD_ERRORS}"
    )


def get_data(key):
    """Returns a copy of the rows of one dataset."""
    return list(_ROWS[key])


def dataset_label(key):
    return DATASETS[key]["label"]


# ============================================================
# COLUMN METADATA
# ============================================================

def get_unique_values(key, column, allowed_departments=None):
    """Sorted unique non-empty values of a column (only rows the user may see)."""

    return sorted(
        {
            str(row.get(column)).strip()
            for row in _visible_rows(key, allowed_departments)
            if row.get(column) not in (None, "")
        }
    )


def resolve_sort_column(key, name):
    """Friendly sort name ('amount', 'date', ...) or a real column -> column."""

    if not name:
        return None

    cfg = DATASETS[key]
    n = _norm(name)

    friendly = {
        "amount": cfg["amount"],
        "budget": cfg["budget"],
        "target": cfg["budget"],
        "date": cfg["date"],
        "id": cfg["id"],
        "status": cfg.get("status_column"),
    }

    if friendly.get(n):
        return friendly[n]

    return resolve_column(cfg["columns"], name)


def resolve_group_column(key, name):

    if not name:
        return None

    if _norm(name) == "month":
        return "Month"

    return resolve_column(DATASETS[key]["group_columns"], name)


def resolve_measure(key, name):

    if not name:
        return None

    cfg = DATASETS[key]
    n = _norm(name)

    if n == "amount":
        return cfg["amount"]

    if n in ("budget", "target"):
        return cfg["budget"]

    return resolve_column(cfg["measures"], name)


# ============================================================
# VALIDATION (no-hallucination layer)
# ============================================================

def match_value(options, value, exact_only=False):
    """
    Maps a user/LLM supplied value to a REAL value in the dataset.
    Returns (canonical_value or None, was_fuzzy_match).
    IDs (exact_only) are never guessed.
    """

    lookup = {option.lower(): option for option in options}
    key = str(value).strip().lower()

    if key in lookup:
        return lookup[key], False

    if exact_only:
        return None, False

    # "Chennai" -> "Chennai WH" (only when exactly one option fits)
    if len(key) >= 3:
        hits = [
            option for option in options
            if key in option.lower()
            or (len(option) >= 3 and option.lower() in key)
        ]
        if len(hits) == 1:
            return hits[0], True

    close = difflib.get_close_matches(key, list(lookup), n=1, cutoff=0.75)

    if close:
        return lookup[close[0]], True

    return None, False


def _as_list(value):
    return value if isinstance(value, (list, tuple)) else [value]


def validate_filters(key, filters, allowed_departments=None):
    """
    Checks every filter against the real data of one dataset.

    Filter structure:
        equals   {column: value or [values]}   categorical columns
        contains {column: text}                partial match on text columns
        exclude  {column: value or [values]}   rows to leave out
        record_id, start_date, end_date, min_amount, max_amount

    Returns (clean_filters, warnings, errors).
    """

    cfg = DATASETS[key]
    filters = filters or {}

    clean = {"equals": {}, "contains": {}, "exclude": {}}
    warnings, errors = [], []

    cat_cols = cfg["filter_columns"]
    text_cols = cfg["text_columns"]

    def check_values(column, values, target):

        options = get_unique_values(key, column, allowed_departments)
        exact_only = column in cfg["exact_columns"]
        matched = []

        for value in _as_list(values):

            hit, fuzzy = match_value(options, value, exact_only)

            if hit is None:
                shown = ", ".join(options[:15])
                more = " ..." if len(options) > 15 else ""
                errors.append(
                    f"'{value}' is not a valid {column.lower()} in the "
                    f"{cfg['label']} data."
                    + (f" Available values: {shown}{more}." if shown else "")
                )
            else:
                if fuzzy:
                    warnings.append(
                        f"I interpreted {column.lower()} '{value}' as '{hit}'."
                    )
                matched.append(hit)

        if matched:
            clean[target].setdefault(column, [])
            for m in matched:
                if m not in clean[target][column]:
                    clean[target][column].append(m)

    def contains_filter(column, text):
        if str(text).strip():
            clean["contains"][column] = str(text).strip()

    # ---- equals ----
    for raw_col, values in (filters.get("equals") or {}).items():

        column = resolve_column(cat_cols, raw_col)

        if column:
            check_values(column, values, "equals")
            continue

        column = resolve_column(text_cols, raw_col)

        if column:                                   # lenient: text column
            for v in _as_list(values):
                contains_filter(column, v)
            continue

        errors.append(
            f"I cannot filter the {cfg['label']} data by '{raw_col}'. "
            f"You can filter by: {', '.join(cat_cols + text_cols)}."
        )

    # ---- contains ----
    for raw_col, text in (filters.get("contains") or {}).items():

        column = resolve_column(text_cols, raw_col)

        if column:
            contains_filter(column, text)
            continue

        column = resolve_column(cat_cols, raw_col)

        if column:                                   # lenient: categorical
            check_values(column, text, "equals")
            continue

        errors.append(
            f"I cannot search the {cfg['label']} data by '{raw_col}'. "
            f"You can filter by: {', '.join(cat_cols + text_cols)}."
        )

    # ---- exclude ----
    for raw_col, values in (filters.get("exclude") or {}).items():

        column = resolve_column(cat_cols, raw_col)

        if column:
            check_values(column, values, "exclude")
        else:
            errors.append(
                f"I cannot exclude by '{raw_col}' in the {cfg['label']} data."
            )

    # ---- record id (never fuzzy matched) ----
    if filters.get("record_id") not in (None, ""):
        clean["record_id"] = str(filters["record_id"]).strip()

    # ---- dates ----
    for name in ("start_date", "end_date"):
        value = filters.get(name)
        if value in (None, ""):
            continue
        try:
            clean[name] = parse_date_input(value).strftime("%Y-%m-%d")
        except ValueError as exc:
            errors.append(str(exc))

    if (
        "start_date" in clean
        and "end_date" in clean
        and clean["start_date"] > clean["end_date"]
    ):
        clean["start_date"], clean["end_date"] = (
            clean["end_date"], clean["start_date"]
        )
        warnings.append(
            "The start date was after the end date, so I swapped them."
        )

    # ---- amounts ----
    for name in ("min_amount", "max_amount"):
        value = filters.get(name)
        if value in (None, ""):
            continue
        try:
            clean[name] = float(value)
        except (TypeError, ValueError):
            errors.append(f"'{value}' is not a valid amount.")

    if (
        "min_amount" in clean
        and "max_amount" in clean
        and clean["min_amount"] > clean["max_amount"]
    ):
        clean["min_amount"], clean["max_amount"] = (
            clean["max_amount"], clean["min_amount"]
        )
        warnings.append(
            "The minimum amount was above the maximum, so I swapped them."
        )

    return clean, warnings, errors


# ============================================================
# SORTING
# ============================================================

def _sort_key(cfg, column):

    if column in cfg["numeric_columns"]:
        return lambda row: _to_float(row.get(column))

    if column in cfg["date_columns"]:
        return lambda row: row.get(column) or datetime.min

    return lambda row: str(row.get(column, "")).lower()


def sort_rows(key, rows, sort_by=None, sort_order=None):
    """Sorts rows by a friendly or real column name."""

    column = resolve_sort_column(key, sort_by)

    if not column:
        return rows

    cfg = DATASETS[key]

    default_desc = (
        column in cfg["numeric_columns"] or column in cfg["date_columns"]
    )

    order = (sort_order or ("desc" if default_desc else "asc")).lower()

    return sorted(
        rows,
        key=_sort_key(cfg, column),
        reverse=(order == "desc")
    )


# ============================================================
# SEARCH / FILTER RECORDS
# ============================================================

def _visible_rows(key, allowed_departments=None):
    """Rows the user may see (department restriction applied)."""

    rows = _ROWS[key]

    if allowed_departments is None:
        return rows

    column = DATASETS[key]["department"]

    if not column:                       # cannot be restricted -> hide all
        return []

    allowed = {d.lower() for d in allowed_departments}

    return [
        row for row in rows
        if str(row.get(column, "")).lower() in allowed
    ]


def search_records(
    key,
    filters=None,
    sort_by=None,
    sort_order=None,
    limit=None,
    allowed_departments=None
):
    """
    Filters, sorts and (optionally) limits the records of one dataset.

    allowed_departments : list of departments the current user may see
                          (None = no restriction).
    """

    cfg = DATASETS[key]
    f = filters or {}

    result = list(_visible_rows(key, allowed_departments))

    # ---- record id ----
    if f.get("record_id"):
        target = str(f["record_id"]).lower()
        result = [
            r for r in result
            if str(r.get(cfg["id"], "")).lower() == target
        ]

    # ---- equals / exclude / contains ----
    for column, values in (f.get("equals") or {}).items():
        wanted = {str(v).lower() for v in _as_list(values)}
        result = [
            r for r in result
            if str(r.get(column, "")).strip().lower() in wanted
        ]

    for column, values in (f.get("exclude") or {}).items():
        banned = {str(v).lower() for v in _as_list(values)}
        result = [
            r for r in result
            if str(r.get(column, "")).strip().lower() not in banned
        ]

    for column, text in (f.get("contains") or {}).items():
        needle = str(text).lower()
        result = [
            r for r in result
            if needle in str(r.get(column, "")).lower()
        ]

    # ---- dates ----
    start = parse_date_input(f.get("start_date"))
    end = parse_date_input(f.get("end_date"), end_of_day=True)

    if start:
        result = [
            r for r in result
            if r.get(cfg["date"]) and r[cfg["date"]] >= start
        ]

    if end:
        result = [
            r for r in result
            if r.get(cfg["date"]) and r[cfg["date"]] <= end
        ]

    # ---- amount range (on the dataset's main amount column) ----
    if f.get("min_amount") is not None:
        result = [
            r for r in result
            if _to_float(r.get(cfg["amount"])) >= float(f["min_amount"])
        ]

    if f.get("max_amount") is not None:
        result = [
            r for r in result
            if _to_float(r.get(cfg["amount"])) <= float(f["max_amount"])
        ]

    result = sort_rows(key, result, sort_by, sort_order)

    if limit:
        result = result[:int(limit)]

    return result


# ============================================================
# BASIC CALCULATIONS
# ============================================================

def calculate_total(rows, column):

    return sum(_to_float(row.get(column)) for row in rows)


def calculate_average(rows, column):

    if not rows:
        return 0.0

    return calculate_total(rows, column) / len(rows)


def count_records(rows):

    return len(rows)


def row_kind(key, row):
    """'income' or 'expense' for one row."""

    cfg = DATASETS[key]

    if cfg["type_column"]:
        return str(row.get(cfg["type_column"], "")).strip().lower()

    return cfg["nature"]


def _primary_rows(key, rows):

    kind = DATASETS[key]["primary_kind"]

    return [row for row in rows if row_kind(key, row) == kind]


def income_expense_summary(key, rows):
    """Total income and total expense (on the dataset's main amount)."""

    amount = DATASETS[key]["amount"]

    total_income = 0.0
    total_expense = 0.0

    for row in rows:

        value = _to_float(row.get(amount))
        kind = row_kind(key, row)

        if kind == "income":
            total_income += value
        elif kind == "expense":
            total_expense += value

    return {
        "income": total_income,
        "expense": total_expense,
        "net": total_income - total_expense
    }


def summarize(key, rows, measure=None):
    """Income, expense, net, count, total and average for a set of rows."""

    cfg = DATASETS[key]
    measure = measure or cfg["amount"]

    base = income_expense_summary(key, rows)
    base["count"] = len(rows)
    base["total"] = calculate_total(rows, measure)
    base["average"] = calculate_average(rows, measure)

    primary = _primary_rows(key, rows)
    actual = calculate_total(primary, cfg["amount"])
    budget = calculate_total(primary, cfg["budget"])

    base["primary_actual"] = actual
    base["primary_budget"] = budget
    base["variance"] = _variance(cfg, budget, actual)

    return base


# ============================================================
# TOP / BOTTOM RECORDS
# ============================================================

def get_top_expenses(key, rows, n=5, measure=None):
    """Highest records (expenses, or orders for sales) by a numeric column."""

    measure = measure or DATASETS[key]["amount"]

    return sorted(
        _primary_rows(key, rows),
        key=lambda row: _to_float(row.get(measure)),
        reverse=True
    )[:n]


def get_bottom_expenses(key, rows, n=5, measure=None):
    """Lowest records by a numeric column."""

    measure = measure or DATASETS[key]["amount"]

    return sorted(
        _primary_rows(key, rows),
        key=lambda row: _to_float(row.get(measure))
    )[:n]


# ============================================================
# BUDGET VS ACTUAL  +  VARIANCE DETECTION
# ============================================================
# Convention: a POSITIVE variance is always FAVOURABLE.
#   cost datasets   (transactions, procurement, payroll):
#                   variance = budget - actual   (positive = under budget)
#   revenue dataset (sales):
#                   variance = actual - target   (positive = above target)
# Budget vs actual is calculated on the dataset's "primary" rows only
# (expenses for transactions, all rows for the others).

def _variance(cfg, budget, actual):

    if cfg["primary_kind"] == "expense":
        return budget - actual

    return actual - budget


def status_labels(key):
    """(unfavourable label, favourable label) for this dataset."""

    if DATASETS[key]["primary_kind"] == "expense":
        return "Over budget", "Under budget"

    return "Below target", "Above target"


def _variance_status(key, budget, actual, threshold_pct):

    cfg = DATASETS[key]

    if budget <= 0:
        return "No budget set" if actual > 0 else "No activity"

    pct = _variance(cfg, budget, actual) / budget * 100
    bad, good = status_labels(key)

    if pct < -threshold_pct:
        return bad

    if pct > threshold_pct:
        return good

    return "Within tolerance"


def budget_vs_actual(key, rows, threshold_pct=10.0):

    cfg = DATASETS[key]
    primary = _primary_rows(key, rows)

    actual = calculate_total(primary, cfg["amount"])
    budget = calculate_total(primary, cfg["budget"])

    variance = _variance(cfg, budget, actual)
    variance_pct = (variance / budget * 100) if budget else None

    return {
        "budget": budget,
        "actual": actual,
        "variance": variance,
        "variance_pct": variance_pct,
        "status": _variance_status(key, budget, actual, threshold_pct)
    }


def _group_key(cfg, row, group_by):

    if group_by == "Month":
        when = row.get(cfg["date"])
        return when.strftime("%Y-%m") if when else "Unknown"

    return str(row.get(group_by) or "Unknown")


def group_summary(key, rows, group_by=None):
    """
    Groups rows and returns one summary row per group.

    transactions : Count, Income, Expense, Net, Budget, Variance
    others       : Count, <actual label>, <budget label>, Variance
    """

    cfg = DATASETS[key]

    group_by = group_by or cfg["default_group"]

    if group_by != "Month" and group_by not in cfg["group_columns"]:
        group_by = cfg["default_group"]

    groups = defaultdict(
        lambda: {"Count": 0, "Income": 0.0, "Expense": 0.0, "Budget": 0.0}
    )

    for row in rows:

        entry = groups[_group_key(cfg, row, group_by)]
        entry["Count"] += 1

        amount = _to_float(row.get(cfg["amount"]))
        kind = row_kind(key, row)

        if kind == "income":
            entry["Income"] += amount
        elif kind == "expense":
            entry["Expense"] += amount

        if kind == cfg["primary_kind"]:
            entry["Budget"] += _to_float(row.get(cfg["budget"]))

    table = []

    for name, e in groups.items():

        actual = e["Expense"] if cfg["primary_kind"] == "expense" else e["Income"]

        if cfg["type_column"]:
            table.append({
                group_by: name,
                "Count": e["Count"],
                "Income": round(e["Income"], 2),
                "Expense": round(e["Expense"], 2),
                "Net": round(e["Income"] - e["Expense"], 2),
                "Budget": round(e["Budget"], 2),
                "Variance": round(_variance(cfg, e["Budget"], actual), 2),
            })
        else:
            table.append({
                group_by: name,
                "Count": e["Count"],
                cfg["amount_label"]: round(actual, 2),
                cfg["budget_label"]: round(e["Budget"], 2),
                "Variance": round(_variance(cfg, e["Budget"], actual), 2),
            })

    sort_key = primary_column(key)

    if group_by == "Month":
        table.sort(key=lambda r: (r[group_by] == "Unknown", r[group_by]))
    else:
        table.sort(key=lambda r: r[sort_key], reverse=True)

    return table


def primary_column(key):
    """Name of the main amount column in group_summary / compare tables."""

    cfg = DATASETS[key]

    if cfg["type_column"]:
        return "Expense" if cfg["primary_kind"] == "expense" else "Income"

    return cfg["amount_label"]


def detect_variances(key, rows, group_by=None, threshold_pct=10.0):
    """
    Compares budget and actual for each group and flags groups whose
    variance is larger than threshold_pct. Worst (most unfavourable)
    group comes first.
    """

    cfg = DATASETS[key]

    group_by = group_by or cfg["default_group"]

    groups = defaultdict(lambda: [0.0, 0.0])      # actual, budget

    for row in _primary_rows(key, rows):
        entry = groups[_group_key(cfg, row, group_by)]
        entry[0] += _to_float(row.get(cfg["amount"]))
        entry[1] += _to_float(row.get(cfg["budget"]))

    table = []

    for name, (actual, budget) in groups.items():

        variance = _variance(cfg, budget, actual)

        table.append({
            group_by: name,
            "Budget": round(budget, 2),
            "Actual": round(actual, 2),
            "Variance": round(variance, 2),
            "Variance %": (
                round(variance / budget * 100, 2) if budget > 0 else None
            ),
            "Status": _variance_status(key, budget, actual, threshold_pct)
        })

    table.sort(
        key=lambda r: (
            r["Variance %"] is None,
            r["Variance %"] if r["Variance %"] is not None else 0
        )
    )

    return table


# ============================================================
# COMPARISON
# ============================================================

def compare_datasets(key, rows_a, rows_b, label_a, label_b):
    """
    Compares two sets of records (two departments, two months, ...).
    Difference = B - A.
    """

    cfg = DATASETS[key]

    summary_a = summarize(key, rows_a)
    summary_b = summarize(key, rows_b)

    if cfg["type_column"]:
        metrics = [
            ("Income", "income"),
            ("Expense", "expense"),
            ("Net", "net"),
            ("Transactions", "count"),
            ("Average Amount", "average"),
        ]
    else:
        metrics = [
            (cfg["amount_label"], "primary_actual"),
            (cfg["budget_label"], "primary_budget"),
            ("Variance", "variance"),
            ("Records", "count"),
            ("Average " + cfg["amount_label"], "average"),
        ]

    table = []

    for name, field in metrics:

        a_value = summary_a[field]
        b_value = summary_b[field]
        difference = b_value - a_value

        percent = (difference / abs(a_value) * 100) if a_value else None

        table.append({
            "Metric": name,
            label_a: round(a_value, 2),
            label_b: round(b_value, 2),
            "Difference": round(difference, 2),
            "% Change": round(percent, 2) if percent is not None else None,
        })

    return {"table": table, "a": summary_a, "b": summary_b}


# ============================================================
# ACCESS CONTROL
# ============================================================
# Passwords come from environment variables (.env). The defaults below
# are for DEMO ONLY - change them before sharing the app.
#
#   datasets        which datasets the user may open (None = all)
#   departments     which departments the user may see (None = all).
#                   A department-restricted user can only open datasets that
#                   have a Department column, because the other datasets
#                   cannot be limited by department.
#   hidden_columns  columns the user can never see, filter by or group by.

_SALT = os.getenv("AUTH_SALT", "ai-pricing-copilot").encode()


def hash_password(password):

    return hashlib.pbkdf2_hmac(
        "sha256",
        str(password).encode(),
        _SALT,
        100_000
    ).hex()


USERS = {
    "admin": {
        "name": "Admin",
        "role": "Admin",
        "datasets": None,
        "departments": None,
        "hidden_columns": [],
        "password_hash": hash_password(os.getenv("ADMIN_PASSWORD", "admin123")),
    },
    "finance": {
        "name": "Finance Manager",
        "role": "Finance Manager",
        "datasets": None,
        "departments": None,
        "hidden_columns": [],
        "password_hash": hash_password(os.getenv("FINANCE_PASSWORD", "finance123")),
    },
    "it_head": {
        "name": "IT Department Head",
        "role": "Department Head",
        "datasets": ["transactions"],
        "departments": ["IT"],          # only sees IT data
        "hidden_columns": [],
        "password_hash": hash_password(os.getenv("ITHEAD_PASSWORD", "ithead123")),
    },
    "analyst": {
        "name": "Analyst",
        "role": "Analyst",
        "datasets": ["transactions", "procurement", "sales"],   # no payroll
        "departments": None,
        # cannot see vendor, supplier or customer names
        "hidden_columns": ["Vendor / Customer", "Supplier", "Customer"],
        "password_hash": hash_password(os.getenv("ANALYST_PASSWORD", "analyst123")),
    },
    "hr_manager": {
        "name": "HR Manager",
        "role": "HR Manager",
        "datasets": ["payroll"],
        "departments": None,
        "hidden_columns": [],
        "password_hash": hash_password(os.getenv("HRMANAGER_PASSWORD", "hrmanager123")),
    },
    "sales_head": {
        "name": "Sales Head",
        "role": "Sales Head",
        "datasets": ["sales"],
        "departments": None,
        "hidden_columns": [],
        "password_hash": hash_password(os.getenv("SALESHEAD_PASSWORD", "saleshead123")),
    },
}


def authenticate(username, password):
    """Returns the user profile (without the hash) or None."""

    record = USERS.get(str(username or "").strip().lower())

    # always hash, so timing does not reveal whether the user exists
    candidate = hash_password(password or "")
    expected = record["password_hash"] if record else hash_password("x")

    if record and hmac.compare_digest(candidate, expected):
        return {
            "username": str(username).strip().lower(),
            "name": record["name"],
            "role": record["role"],
            "datasets": record["datasets"],
            "departments": record["departments"],
            "hidden_columns": list(record["hidden_columns"]),
        }

    return None


def allowed_departments(user):
    """None = unrestricted, otherwise the list of permitted departments."""

    if not user:
        return None

    return user.get("departments")


def dataset_access_problem(user, key):
    """Returns a friendly message if the user may not open this dataset."""

    label = DATASETS[key]["label"]

    permitted = (user or {}).get("datasets")

    if permitted is not None and key not in permitted:
        return f"You do not have permission to view the {label} data."

    departments = allowed_departments(user)

    if departments is not None and not DATASETS[key]["department"]:
        return (
            f"You do not have permission to view the {label} data, because "
            "it cannot be limited to your department."
        )

    return None


def allowed_datasets(user):
    """Keys of the loaded datasets this user may open."""

    return [
        key for key in DATASETS
        if dataset_access_problem(user, key) is None
    ]


def hidden_columns(user):
    """Normalised set of columns the user may not see."""

    return {_norm(c) for c in ((user or {}).get("hidden_columns") or [])}


def sanitize_rows(rows, user):
    """Removes columns the user is not allowed to see."""

    hidden = hidden_columns(user)

    if not hidden:
        return rows

    return [
        {k: v for k, v in row.items() if _norm(k) not in hidden}
        for row in rows
    ]
