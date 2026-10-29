"""Create the SQLite schema and optionally import the supplied workbooks."""

import argparse
import os
import sqlite3
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from backend import database
from backend.data_engine import DATASET_CONFIGS

DATE_FORMATS = (
    "%Y-%m-%d", "%d-%m-%Y", "%m-%d-%Y", "%d/%m/%Y",
    "%m/%d/%Y", "%Y/%m/%d", "%d-%b-%Y",
)


def _parse_date(value, filename, row_number, column):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    for date_format in DATE_FORMATS:
        try:
            return datetime.strptime(str(value).strip(), date_format).date()
        except ValueError:
            continue
    raise ValueError(
        f"{filename}, row {row_number}: invalid date in '{column}'."
    )


def _load_source_rows(dataset, base_dir):
    config = DATASET_CONFIGS[dataset]
    filename = config["file"]
    path = os.path.join(base_dir, "Data", filename)
    if not os.path.isfile(path):
        raise FileNotFoundError(
            f"Source workbook '{path}' was not found. No data was imported."
        )

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = (
            workbook[config["sheet"]]
            if config["sheet"] and config["sheet"] in workbook.sheetnames
            else workbook.active
        )
        headers = [
            str(cell.value).strip() if cell.value is not None else ""
            for cell in sheet[1]
        ]
        expected = database.DATASET_COLUMNS[dataset]
        if headers != expected:
            raise ValueError(
                f"Workbook '{filename}' does not match the approved "
                f"{dataset} schema. Expected columns in order: "
                + ", ".join(expected)
            )

        numeric_columns = set(config["numeric_columns"])
        date_columns = set(config["date_columns"])
        records = []

        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            if not any(value not in (None, "") for value in values):
                continue

            record = dict(zip(headers, values))
            for column in headers:
                value = record[column]
                if isinstance(value, str):
                    value = value.strip() or None

                if column in date_columns:
                    value = _parse_date(value, filename, row_number, column)
                elif column in numeric_columns:
                    if value in (None, ""):
                        value = None
                    else:
                        try:
                            value = Decimal(str(value))
                        except (InvalidOperation, TypeError, ValueError) as exc:
                            raise ValueError(
                                f"{filename}, row {row_number}: invalid "
                                f"number in '{column}'."
                            ) from exc
                        if not value.is_finite():
                            raise ValueError(
                                f"{filename}, row {row_number}: invalid "
                                f"number in '{column}'."
                            )
                record[column] = value

            for column, formula in config["fallbacks"].items():
                if record.get(column) is None:
                    try:
                        value = formula(record)
                        record[column] = (
                            Decimal(str(value)) if value is not None else None
                        )
                    except (KeyError, TypeError, ValueError, InvalidOperation) as exc:
                        raise ValueError(
                            f"{filename}, row {row_number}: could not "
                            f"calculate '{column}'."
                        ) from exc
                    if (
                        record[column] is not None
                        and not record[column].is_finite()
                    ):
                        raise ValueError(
                            f"{filename}, row {row_number}: invalid "
                            f"calculated value for '{column}'."
                        )
            records.append(record)

        return records
    finally:
        workbook.close()


def _column_type(dataset, column):
    config = DATASET_CONFIGS[dataset]
    numeric = set(config["numeric_columns"])
    dates = set(config["date_columns"])
    if column in numeric:
        return "NUMERIC"
    if column in dates:
        return "DATE"
    return "TEXT"


def _table_ddl(dataset):
    declarations = ['"id" INTEGER PRIMARY KEY AUTOINCREMENT']
    for column in database.DATASET_COLUMNS[dataset]:
        declarations.append(
            f"{database._quote_identifier(column)} "
            f"{_column_type(dataset, column)}"
        )
    table = database._quote_identifier(database.TABLES[dataset])
    return "CREATE TABLE IF NOT EXISTS " + table + " (" + ", ".join(
        declarations
    ) + ")"


def _document_ddl():
    return [
        """
        CREATE TABLE IF NOT EXISTS documents (
            id TEXT PRIMARY KEY,
            original_name TEXT NOT NULL,
            storage_name TEXT NOT NULL UNIQUE,
            content_type TEXT NOT NULL,
            size_bytes BIGINT NOT NULL,
            owner TEXT NOT NULL,
            department TEXT,
            status TEXT NOT NULL CHECK (
                status IN ('processing', 'ready', 'failed')
            ),
            ingestion_error TEXT,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS document_chunks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            document_id TEXT NOT NULL REFERENCES documents(id)
                ON DELETE CASCADE,
            chunk_number INTEGER NOT NULL,
            page_number INTEGER,
            section TEXT,
            chunk_text TEXT NOT NULL,
            embedding TEXT NOT NULL,
            UNIQUE (document_id, chunk_number)
        )
        """,
        """
        CREATE VIRTUAL TABLE IF NOT EXISTS document_chunks_fts
        USING fts5(chunk_text, content='document_chunks', content_rowid='id')
        """,
        """
        CREATE TRIGGER IF NOT EXISTS document_chunks_fts_insert
        AFTER INSERT ON document_chunks BEGIN
            INSERT INTO document_chunks_fts(rowid, chunk_text)
            VALUES (new.id, new.chunk_text);
        END
        """,
        """
        CREATE TRIGGER IF NOT EXISTS document_chunks_fts_delete
        AFTER DELETE ON document_chunks BEGIN
            INSERT INTO document_chunks_fts(document_chunks_fts, rowid, chunk_text)
            VALUES ('delete', old.id, old.chunk_text);
        END
        """,
        """
        CREATE TRIGGER IF NOT EXISTS document_chunks_fts_update
        AFTER UPDATE OF chunk_text ON document_chunks BEGIN
            INSERT INTO document_chunks_fts(document_chunks_fts, rowid, chunk_text)
            VALUES ('delete', old.id, old.chunk_text);
            INSERT INTO document_chunks_fts(rowid, chunk_text)
            VALUES (new.id, new.chunk_text);
        END
        """,
        """
        CREATE INDEX IF NOT EXISTS documents_owner_department_idx
        ON documents (owner, department, status)
        """,
        """
        CREATE INDEX IF NOT EXISTS document_chunks_document_idx
        ON document_chunks (document_id, chunk_number)
        """,
    ]


def initialize(seed_from_excel=False):
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    source_rows = {}

    if seed_from_excel:
        source_rows = {
            dataset: _load_source_rows(dataset, base_dir)
            for dataset in database.TABLES
        }

    with database.connect(read_only=False) as connection:
        with connection.cursor() as cursor:
            for dataset in database.TABLES:
                cursor.execute(_table_ddl(dataset))
                table = database._quote_identifier(database.TABLES[dataset])
                cursor.execute("PRAGMA table_info(" + table + ")")
                existing_columns = {row[1] for row in cursor.fetchall()}
                for column in database.DATASET_COLUMNS[dataset]:
                    if column not in existing_columns:
                        cursor.execute(
                            f"ALTER TABLE {table} ADD COLUMN "
                            f"{database._quote_identifier(column)} "
                            f"{_column_type(dataset, column)}"
                        )

                if not seed_from_excel:
                    continue

                columns = database.DATASET_COLUMNS[dataset]
                quoted_columns = ", ".join(
                    database._quote_identifier(column) for column in columns
                )
                placeholders = ", ".join(["?"] * len(columns))

                cursor.execute(f"DELETE FROM {table}")
                cursor.executemany(
                    f"INSERT INTO {table} ({quoted_columns}) "
                    f"VALUES ({placeholders})",
                    [
                        tuple(
                            value.isoformat()
                            if isinstance(value, (date, datetime))
                            else str(value)
                            if isinstance(value, Decimal)
                            else value
                            for value in (record[column] for column in columns)
                        )
                        for record in source_rows[dataset]
                    ],
                )

            for statement in _document_ddl():
                cursor.execute(statement)

    for dataset in database.TABLES:
        record_count = len(source_rows[dataset]) if seed_from_excel else None
        if record_count is None:
            print(f"Initialized schema for {dataset}.")
        else:
            print(f"Imported {record_count} {dataset} records.")
    print("Initialized PDF/Word document schema with SQLite FTS5.")


def main():
    parser = argparse.ArgumentParser(
        description="Initialize the SQLite database for the chatbot."
    )
    parser.add_argument(
        "--seed-from-excel",
        action="store_true",
        help=(
            "Explicitly import the current Data/*.xlsx workbooks. "
            "This replaces existing rows in the four approved dataset tables."
        ),
    )
    args = parser.parse_args()

    initialize(seed_from_excel=args.seed_from_excel)


if __name__ == "__main__":
    main()
